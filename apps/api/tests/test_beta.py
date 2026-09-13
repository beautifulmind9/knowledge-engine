import io
import json
import os
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace
from zipfile import ZipFile
import pytest
from app.db import mock_data as db
from app.db.persistence import STORAGE_ROOT
from app.services.knowledge_extraction import persist_state, consolidate_knowledge_assets


def brief(source_id,mode='writing'):
    return {'situation':'Help participants retain learning','goal':'Use practice breaks and specific goals',
            'output_type':mode,'source_ids':[source_id]}


def output(asset_id,mode='writing'):
    return {'title':'Practical learning','output_type':mode,
        'content': {
            'writing': 'Learning sticks when people get to practice. Build in practice breaks, then invite participants to share what they will try next.',
            'decision_brief': 'Decision\nChoose a learning format.\nOptions\nContinuous presentation or practice with breaks.\nTrade-offs\nPractice leaves less time for lecture.\nRecommendation\nUse practice breaks.\nUncertainties\nCheck prior experience.\nNext steps\nTry one session.',
            'study_guide': 'Learning goals\nExplain the purpose of practice breaks.\nKey ideas\nPractice breaks support retention.\nPractice\nPlan a pause in a learning session.\nReview questions\nWhat helps participants retain learning?\nAnswer guidance\nPractice with breaks.',
            'product_messaging': 'Audience\nFacilitators.\nValue proposition\nMake space for practical learning.\nSupporting messages\nPractice breaks help participants retain learning.\nDraft copy\nGive learners a chance to practice and pause.',
            'playbook': 'Purpose\nSupport retention.\nPrerequisites\nA learning task.\nSteps\n1. Introduce the task.\n2. Include practice breaks.\nChecks\nAsk learners to recall the task.\nExceptions\nAdapt to the audience.',
            'workshop_plan': 'Learning goals\nPractice one useful skill.\nAgenda\nIntroduction, guided practice, break, and reflection.\nBreak\nPause between practice rounds.\nActivities\nTry the skill and explain the result.',
        }.get(mode, 'A sufficiently long unsupported-mode fixture for validation.'),
        'applied_knowledge':[{'asset_id':asset_id,'usage_note':'Uses practice breaks to support retention.'}],
        'design_choices':['The proposed sequence is a drafting choice.']}


def save_manual(client,knowledge,mode='writing'):
    r=client.post('/outputs',json={'brief':brief(knowledge['source_id'],mode),'output':output(knowledge['id'],mode)})
    assert r.status_code==200,r.text
    return r.json()


def test_full_workflow_and_fresh_process(client,knowledge):
    first=save_manual(client,knowledge)
    second=client.post(f"/outputs/{first['id']}/revise",json={'instruction':'Make concise','content':'Take practice breaks to improve retention of the learning.'}).json()
    assert second['parent_output_id']==first['id'] and second['version']==2
    assert client.get(f"/outputs/{first['id']}").json()==first
    assert len(client.get(f"/outputs/{first['id']}/history").json()['items'])==2
    assert 'content' in client.get(f"/outputs/{first['id']}/compare/{second['id']}").json()['changed_fields']
    exported=client.get(f"/outputs/{second['id']}/export")
    assert '## Provenance' in exported.text and knowledge['chunk_id'] in exported.text
    code='from app.db.mock_data import outputs; import json; print(json.dumps(outputs))'
    env={**os.environ,'PYTHONPATH':str(Path(__file__).resolve().parents[1])}
    recovered=json.loads(subprocess.check_output([sys.executable,'-c',code],env=env,cwd='/tmp',text=True))
    assert recovered[0]==first and recovered[1]==second
    assert client.get('/outputs?output_type=writing').json()['count']==2
    assert client.get('/outputs?created_after=2099-01-01T00:00:00Z').json()['count']==0

@pytest.mark.parametrize('mode',['writing','decision_brief','study_guide','product_messaging','playbook','workshop_plan'])
def test_generate_save_revise_modes_single_call(client,knowledge,monkeypatch,mode):
    from app.services import ai_gateway
    calls=[]
    def generate(model,payload,schema):
        calls.append(payload)
        return SimpleNamespace(output_text=json.dumps(output(knowledge['id'],mode)),id='fixture-interaction')
    monkeypatch.setattr(ai_gateway,'generate',generate)
    r=client.post('/workshops/generate',json=brief(knowledge['source_id'],mode))
    assert r.status_code==200,r.text
    saved=r.json()['saved_output']
    assert len(calls)==1 and saved['output_type']==mode
    assert saved['generation_metadata']['quality']['has_provenance']
    assert all(c['passed'] for c in saved['generation_metadata']['quality']['structure_checks'])
    r=client.post(f"/outputs/{saved['id']}/revise",json={'instruction':'Adapt for beginners'})
    assert r.status_code==200,r.text
    assert len(calls)==2 and calls[-1]['revision']['previous_output']==saved['content']
    assert r.json()['version']==2


def test_atomic_import_and_intentional_reprocessing(client,source,asset):
    s,c=source
    job=client.post('/knowledge-extractions',json={'source_id':s['id'],'chunk_id':c['id']}).json()['job']
    invalid={**asset,'source_id':'wrong_source'}
    r=client.post(f"/knowledge-extractions/{job['id']}/results",json={'assets':[asset,invalid]})
    assert r.status_code==400 and not db.knowledge_assets
    assert client.post(f"/knowledge-extractions/{job['id']}/results",json={'assets':[asset]}).status_code==200
    original=dict(db.knowledge_assets[0])
    assert client.post('/knowledge-extractions',json={'source_id':s['id'],'chunk_id':c['id']}).status_code==400
    nextjob=client.post('/knowledge-extractions',json={'source_id':s['id'],'chunk_id':c['id'],'reprocess':True}).json()['job']
    assert nextjob['extraction_version']==2
    assert client.post(f"/knowledge-extractions/{nextjob['id']}/results",json={'assets':[invalid]}).status_code==400
    assert db.knowledge_assets[0]==original
    assert client.post(f"/knowledge-extractions/{nextjob['id']}/results",json={'assets':[asset]}).status_code==200
    assert len(db.knowledge_assets)==2 and not db.knowledge_assets[0]['active']
    assert client.get('/knowledge-assets').json()['count']==1
    audit=client.get(f"/sources/{s['id']}/audit").json()
    assert audit['superseded_assets']==[original['id']] and audit['confidence_distribution']=={'4':1}
    assert db.knowledge_assets[1]['chapter_or_section']=='Focus'


def test_invalid_ai_outputs_never_saved(client,knowledge,monkeypatch):
    from app.services import ai_gateway
    for raw in ('{}','not JSON',json.dumps(output('invented_asset')),json.dumps(output(knowledge['id'],'other_mode'))):
        monkeypatch.setattr(ai_gateway,'generate',lambda *args:SimpleNamespace(output_text=raw))
        response=client.post('/workshops/generate',json=brief(knowledge['source_id']))
        assert response.status_code in (400,422)
        assert not db.outputs


def test_missing_key_and_empty_retrieval(client,knowledge,monkeypatch):
    monkeypatch.delenv('GEMINI_API_KEY',raising=False)
    assert client.post('/workshops/generate',json=brief(knowledge['source_id'])).status_code==503
    assert not db.outputs
    assert client.post('/workshops/prepare',json=brief('unknown')).status_code==404
    r=client.post('/workshops/generate',json={**brief(knowledge['source_id']),'situation':'zzzz','goal':'zzzz'})
    assert r.status_code==400


def test_stale_and_failed_recovery(client,source):
    s,c=source
    j=client.post('/knowledge-extractions',json={'source_id':s['id'],'chunk_id':c['id']}).json()['job']
    stored=db.extraction_jobs[0]
    stored.update(status='running',started_at='2000-01-01T00:00:00+00:00')
    assert client.post(f"/knowledge-extractions/{j['id']}/recover").json()['status']=='pending_ai'
    from datetime import datetime,timezone
    stored.update(status='running',started_at=datetime.now(timezone.utc).isoformat())
    assert client.post(f"/knowledge-extractions/{j['id']}/recover").status_code==400
    stored['status']='failed'
    assert client.post(f"/knowledge-extractions/{j['id']}/recover").status_code==200


def test_quota_stop_no_retry_and_resume(client,knowledge,monkeypatch):
    from app.services import ai_gateway
    monkeypatch.setenv('GEMINI_API_KEY','fixture-key')
    monkeypatch.setenv('GEMINI_FREE_TIER_CONFIRMED','true')
    calls=[]
    class QuotaError(Exception):
        code=429
    class FakeClient:
        def __init__(self,**kwargs):
            assert kwargs['http_options']['retry_options']['attempts']==1
            self.interactions=self
            self.sdk_configuration=SimpleNamespace(retry_config=SimpleNamespace(strategy='attempt-count-backoff'))
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def create(self,**kwargs):
            calls.append(kwargs)
            raise QuotaError('quota exceeded')
    monkeypatch.setattr(ai_gateway.genai,'Client',FakeClient)
    assert client.post('/workshops/generate',json=brief(knowledge['source_id'])).status_code==429
    assert client.get('/usage').json()['paused']
    assert client.post('/workshops/generate',json=brief(knowledge['source_id'])).status_code==429
    assert len(calls)==1
    assert not client.post('/usage/resume').json()['paused']
    assert client.get('/usage').json()['calls_today']==1


def test_daily_cap_and_free_confirmation(client,knowledge,monkeypatch):
    monkeypatch.setenv('GEMINI_API_KEY','fixture-key')
    assert client.post('/workshops/generate',json=brief(knowledge['source_id'])).status_code==503
    monkeypatch.setenv('GEMINI_FREE_TIER_CONFIRMED','true')
    monkeypatch.setenv('GEMINI_DAILY_CALL_LIMIT','0')
    assert client.post('/workshops/generate',json=brief(knowledge['source_id'])).status_code==429
    assert not db.outputs


def test_multi_source_agreements_tensions_and_scope(client,knowledge):
    source=client.post('/sources',json={'title':'Alternative notes','library_id':db.sources[0]['library_id']}).json()
    common={**knowledge,'id':'asset_other','source_id':source['id'],'chunk_id':'chunk_other'}
    opposing={**knowledge,'id':'asset_tension','source_id':source['id'],'chunk_id':'chunk_other',
        'what_it_says':'Do not use practice breaks to help participants retain learning.',
        'evidence':'Do not use practice breaks to help participants retain learning.'}
    db.knowledge_assets.extend([common,opposing])
    groups=consolidate_knowledge_assets(db.knowledge_assets)
    assert len(groups)==2
    shared=next(g for g in groups if g['support_count']==2)
    assert len(shared['source_ids'])==2 and len(shared['evidence_trail'])==2
    r=client.post('/workshops/prepare',json={**brief(knowledge['source_id']),'source_ids':[], 'library_id':db.sources[0]['library_id']}).json()
    assert r['synthesis_context']['agreements'] and r['synthesis_context']['tensions']
    empty=client.post('/libraries',json={'name':'Empty'}).json()
    assert client.post('/workshops/prepare',json={**brief(knowledge['source_id']),'source_ids':[],'library_id':empty['id']}).json()['knowledge_unit_count']==0
    assert client.post('/workshops/prepare',json={**brief(knowledge['source_id']),'library_id':empty['id']}).status_code==400
    assert client.post('/workshops/prepare',json={**brief(knowledge['source_id']),'asset_ids':['invented']}).status_code==404


def test_deletion_backup_and_integrity(client,knowledge):
    saved=save_manual(client,knowledge)
    sid=knowledge['source_id']
    assert client.get('/data/integrity').json()['ok']
    archive=ZipFile(io.BytesIO(client.get('/data/export').content))
    backup=json.loads(archive.read('storage/state.json'))
    assert backup['outputs'][0]['id']==saved['id']
    assert any(n.endswith('.md') for n in archive.namelist())
    assert client.delete(f'/sources/{sid}').status_code==409
    assert client.delete(f"/outputs/{saved['id']}").status_code==409
    assert client.delete(f'/sources/{sid}?delete_outputs=true').status_code==200
    assert not db.sources and not db.knowledge_assets and not db.extraction_jobs and not db.outputs
    assert client.get('/data/integrity').json()['ok']


def test_output_delete_history(client,knowledge):
    saved=save_manual(client,knowledge)
    assert client.delete(f"/outputs/{saved['id']}?entire_history=true").status_code==200
    assert client.get(f"/outputs/{saved['id']}").status_code==404
    assert db.sources and db.knowledge_assets


def test_invalid_upload_empty_file_and_missing_source(client,source):
    s,c=source
    assert client.post('/sources',json={'title':'Bad','library_id':'missing'}).status_code==404
    assert client.post('/sources/missing/upload',files={'file':('n.txt','hello')}).status_code==404
    assert client.post(f"/sources/{s['id']}/upload",files={'file':('bad.exe','x')}).status_code==400
    assert client.post(f"/sources/{s['id']}/upload",files={'file':('empty.txt','')}).status_code==400
    assert client.post(f"/sources/{s['id']}/upload",files={'file':('blank.txt','   ')}).status_code==200
    assert client.post(f"/sources/{s['id']}/process").status_code==400

@pytest.mark.parametrize('filename,content',[('notes.txt','Use practice breaks.'),('notes.html','<h1>Learning</h1><p>Use practice breaks.</p><script>bad()</script>')])
def test_cross_format_pipeline(client,source,filename,content):
    s,c=source
    assert client.post(f"/sources/{s['id']}/upload",files={'file':(filename,content)}).status_code==200
    assert client.post(f"/sources/{s['id']}/process").status_code==200
    text=client.get(f"/sources/{s['id']}/extracted-text").text
    assert 'Use practice breaks.' in text and 'bad()' not in text


def test_private_local_web_and_cross_origin(client):
    assert 'Knowledge Engine' in client.get('/').text
    assert client.get('/static/app.js').status_code==200
    assert client.get('/storage/state.json').status_code==404
    assert client.post('/libraries',json={'name':'evil'},headers={'Origin':'https://evil.example'}).status_code==403
    assert client.get('/libraries',headers={'Host':'evil.example'}).status_code==400
    assert not db.libraries


def test_missing_chunk_file_not_silently_complete(client,source):
    s,c=source
    Path(db.sources[0]['chunks_path']).unlink()
    assert client.get(f"/sources/{s['id']}/interpretation").status_code==404
    assert not client.get('/data/integrity').json()['ok']


def test_legacy_migration_and_corruption(tmp_path):
    legacy={'libraries':[],'sources':[],'extraction_jobs':[],'knowledge_assets':[]}
    storage=tmp_path/'storage';storage.mkdir()
    path=storage/'state.json';path.write_text(json.dumps(legacy))
    env={**os.environ,'KNOWLEDGE_ENGINE_STORAGE':str(storage),'PYTHONPATH':str(Path(__file__).resolve().parents[1])}
    result=subprocess.run([sys.executable,'-c','from app.db.persistence import load_state; import json; print(json.dumps(load_state()))'],env=env,capture_output=True,text=True)
    assert result.returncode==0 and json.loads(result.stdout)==legacy
    assert json.loads(path.read_text())==legacy
    (storage/'knowledge.sqlite3').unlink();path.write_text('{bad json')
    result=subprocess.run([sys.executable,'-c','from app.db.persistence import load_state; load_state()'],env=env,capture_output=True,text=True)
    assert result.returncode!=0 and path.read_text()=='{bad json'


def test_revision_preserves_tone_limit_and_synthesis(client,knowledge,monkeypatch):
    from app.services import ai_gateway
    requests=[]
    def generate(model,payload,schema):
        requests.append(payload)
        return SimpleNamespace(output_text=json.dumps(output(knowledge['id'])),id='fixture')
    monkeypatch.setattr(ai_gateway,'generate',generate)
    payload={**brief(knowledge['source_id']),'tone_or_style':'Warm and concise','limit':3}
    original=client.post('/workshops/generate',json=payload).json()['saved_output']
    revised=client.post(f"/outputs/{original['id']}/revise",json={'instruction':'Adapt to a beginner'}).json()
    assert revised['brief']['tone_or_style']=='Warm and concise'
    assert revised['brief']['limit']==3
    assert requests[1]['output_preferences']==requests[0]['output_preferences']
    assert requests[1]['synthesis_context']==requests[0]['synthesis_context']


def test_search_stays_in_library(client,knowledge):
    other=client.post('/libraries',json={'name':'Unrelated'}).json()
    for consolidated in ('true','false'):
        assert client.get(f"/knowledge-assets/search?q=practice&library_id={other['id']}&consolidated={consolidated}").json()['count']==0
        assert client.get(f"/knowledge-assets/search?q=practice&library_id={db.sources[0]['library_id']}&consolidated={consolidated}").json()['count']==1
    assert client.get('/knowledge-assets/search?q=practice&library_id=missing').status_code==404
    assert client.get(f"/knowledge-assets/search?q=practice&library_id={other['id']}&source_id={knowledge['source_id']}").status_code==400


def test_manual_interpretation_updates_source_status(client,knowledge):
    assert client.get(f"/sources/{knowledge['source_id']}").json()['processing_status']=='knowledge_interpreted'


@pytest.mark.parametrize('count,expected',[(1000,1),(1001,1),(1200,1),(1201,2),(2200,2),(2201,3)])
def test_chunk_windows_cover_without_redundant_tail(count,expected):
    from app.services.text_chunking import chunk_text
    chunks=chunk_text('fixture',' '.join(str(i) for i in range(count)))
    assert len(chunks)==expected
    assert chunks[-1]['end_word_index']==count
    assert {word for c in chunks for word in c['text'].split()}=={str(i) for i in range(count)}


def test_backup_restore_relocates_and_preserves_history(client,knowledge,tmp_path):
    saved=save_manual(client,knowledge)
    archive=ZipFile(io.BytesIO(client.get('/data/export').content))
    archive.extractall(tmp_path)
    restored=tmp_path/'custom-vault'
    (tmp_path/'storage').rename(restored)
    env={**os.environ,'KNOWLEDGE_ENGINE_STORAGE':str(restored),'PYTHONPATH':str(Path(__file__).resolve().parents[1])}
    code='from app.routers.control import integrity; from app.db.mock_data import outputs; import json; print(json.dumps({"audit":integrity(),"outputs":outputs}))'
    result=json.loads(subprocess.check_output([sys.executable,'-c',code],env=env,text=True,cwd='/tmp'))
    assert result['audit']['ok'],result
    assert result['outputs'][0]==saved


def test_replacing_uninterpreted_upload_removes_old_files(client,source):
    s,c=source
    old=[Path(db.sources[0][key]) for key in ('file_path','extracted_text_path','chunks_path')]
    response=client.post(f"/sources/{s['id']}/upload",files={'file':('replacement.txt','Different original notes.')})
    assert response.status_code==200
    assert all(not path.exists() for path in old)
    assert client.get('/data/integrity').json()['ok']


def test_two_source_generation_keeps_material_provenance(client,knowledge,monkeypatch):
    from app.services import ai_gateway
    source=client.post('/sources',json={'library_id':db.sources[0]['library_id'],'title':'Understanding checks'}).json()
    sid=source['id']
    evidence='Ask participants to explain their next action to check understanding.'
    client.post(f'/sources/{sid}/upload',files={'file':('checks.txt',evidence)})
    client.post(f'/sources/{sid}/process')
    chunk=client.get(f'/sources/{sid}/chunks').json()['items'][0]
    job=client.post('/knowledge-extractions',json={'source_id':sid,'chunk_id':chunk['id']}).json()['job']
    asset={'asset_type':'principle','confidence_score':4,'source_id':sid,'chunk_id':chunk['id'],'title':'Check understanding','what_it_says':evidence,'evidence':evidence,'keywords':['understanding','learning']}
    second=client.post(f"/knowledge-extractions/{job['id']}/results",json={'assets':[asset]}).json()['assets'][0]
    generated=output(knowledge['id'])
    generated['applied_knowledge'].append({'asset_id':second['id'],'usage_note':'Adds an understanding check after practice.'})
    seen=[]
    def generate(model,payload,schema):
        seen.append(payload)
        return SimpleNamespace(output_text=json.dumps(generated))
    monkeypatch.setattr(ai_gateway,'generate',generate)
    saved=client.post('/workshops/generate',json={**brief(knowledge['source_id']),'source_ids':[knowledge['source_id'],sid]}).json()['saved_output']
    assert set(saved['source_ids'])=={knowledge['source_id'],sid}
    assert len(seen)==1 and len(seen[0]['knowledge_units'])==2
    assert len(saved['applied_asset_ids'])==2


def test_live_demo_seed_and_restart(tmp_path):
    import socket
    import time
    import urllib.request
    root=Path(__file__).resolve().parents[3]
    env={**os.environ,'KNOWLEDGE_ENGINE_STORAGE':str(tmp_path/'demo-vault')}
    env.pop('GEMINI_API_KEY',None)
    with socket.socket() as sock:
        sock.bind(('127.0.0.1',0))
        port=sock.getsockname()[1]
    url=f'http://127.0.0.1:{port}'
    opener=urllib.request.build_opener(urllib.request.ProxyHandler({}))
    def read(path):
        with opener.open(url+path,timeout=1) as response:
            return json.load(response)
    def start():
        process=subprocess.Popen([sys.executable,str(root/'scripts/run.py'),'--port',str(port)],env=env,stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
        for _ in range(100):
            if process.poll() is not None:
                raise AssertionError(process.stderr.read().decode())
            try:
                read('/libraries')
                return process
            except OSError:
                time.sleep(.05)
        process.terminate();process.wait(timeout=5)
        raise AssertionError('Server did not become ready')
    process=start()
    try:
        command=[sys.executable,str(root/'scripts/seed_demo.py'),'--url',url]
        seeded=json.loads(subprocess.check_output(command,env=env,text=True))
        assert seeded['versions']==2 and seeded['ai_calls']==0 and seeded['agreements']>=1
        assert len(read('/outputs')['items'])==2
        again=json.loads(subprocess.check_output(command,env=env,text=True))
        assert again['status'].startswith('Demo already exists')
        assert read('/data/integrity')['ok']
        original=read('/outputs/'+seeded['output_id'])
    finally:
        process.terminate();process.wait(timeout=5)
    process=start()
    try:
        assert read('/outputs/'+seeded['output_id'])==original
        assert read('/usage')['calls_today']==0
    finally:
        process.terminate();process.wait(timeout=5)


@pytest.mark.parametrize('provider_status,expected', [(200,200),(429,429),(503,502),(0,502)])
def test_real_sdk_serialization_without_network(client,knowledge,monkeypatch,provider_status,expected):
    import httpx
    from google import genai
    from app.services import ai_gateway
    original_client=genai.Client
    requests=[]
    def transport(request):
        requests.append(json.loads(request.content))
        if provider_status == 0:
            raise httpx.ReadTimeout('fixture timeout',request=request)
        if provider_status != 200:
            return httpx.Response(provider_status,json={'error': {'code':provider_status,'message':'fixture provider failure'}})
        return httpx.Response(200,json={'id':'sdk-fixture','status':'completed','steps':[{'type':'model_output','content':[{'type':'text','text':json.dumps(output(knowledge['id']))}]}]})
    def factory(**kwargs):
        kwargs['http_options'].update(client_args={'transport':httpx.MockTransport(transport),'trust_env':False},async_client_args={'trust_env':False})
        return original_client(**kwargs)
    monkeypatch.setattr(ai_gateway.genai,'Client',factory)
    monkeypatch.setenv('GEMINI_API_KEY','fixture-key')
    monkeypatch.setenv('GEMINI_FREE_TIER_CONFIRMED','true')
    response=client.post('/workshops/generate',json=brief(knowledge['source_id']))
    assert response.status_code==expected,response.text
    assert len(requests)==1
    assert requests[0]['store'] is False
    assert requests[0]['response_format']['mime_type']=='application/json'
    if expected == 200:
        assert response.json()['saved_output']['generation_metadata']['model_response_id']=='sdk-fixture'
    else:
        assert not db.outputs


def test_docx_preserves_heading_and_table_order(tmp_path):
    from docx import Document
    from app.services.text_extraction import extract_text_from_file
    document=Document()
    document.add_heading('Learning goals',level=1)
    document.add_paragraph('Practice one useful skill.')
    document.add_table(rows=1,cols=1).cell(0,0).text='Check understanding.'
    document.add_paragraph('Close with reflection.')
    path=tmp_path/'original.docx';document.save(path)
    text=extract_text_from_file(str(path),'.docx')
    assert text.startswith('# Learning goals')
    assert text.index('Practice one') < text.index('Check understanding') < text.index('Close with')


def test_epub_uses_reading_order_and_headings(tmp_path):
    from ebooklib import epub
    from app.services.text_extraction import extract_text_from_file
    book=epub.EpubBook();book.set_identifier('synthetic');book.set_title('Original learning notes');book.set_language('en')
    first=epub.EpubHtml(title='First',file_name='first.xhtml',lang='en')
    first.content='<h1>Learning goals</h1><p>Choose one useful skill.</p>'
    second=epub.EpubHtml(title='Second',file_name='second.xhtml',lang='en')
    second.content='<h1>Practice</h1><p>Try the skill with breaks.</p>'
    book.add_item(second);book.add_item(first)
    book.spine=[first,second]
    book.add_item(epub.EpubNav());book.add_item(epub.EpubNcx())
    path=tmp_path/'original.epub';epub.write_epub(str(path),book)
    text=extract_text_from_file(str(path),'.epub')
    assert text.index('# Learning goals') < text.index('# Practice')


def test_daily_budget_is_visible_and_cannot_be_resumed(client,monkeypatch):
    monkeypatch.setenv('GEMINI_DAILY_CALL_LIMIT','0')
    status=client.get('/usage').json()
    assert status['paused'] and status['daily_limit_reached']
    assert client.post('/usage/resume').json()['paused']


def test_explicit_unit_selection_retains_supporting_sources(client,knowledge):
    source=client.post('/sources',json={'library_id':db.sources[0]['library_id'],'title':'Additional supporting notes'}).json()
    db.knowledge_assets.append({**knowledge,'id':'asset_support','source_id':source['id'],'chunk_id':'chunk_support'})
    prepared=client.post('/workshops/prepare',json={**brief(knowledge['source_id']),
        'source_ids':[knowledge['source_id'],source['id']], 'asset_ids':[knowledge['id']]}).json()
    assert prepared['knowledge_unit_count']==1
    unit=prepared['knowledge_units'][0]
    assert set(unit['asset_ids'])=={knowledge['id'],'asset_support'}
    assert set(unit['source_ids'])=={knowledge['source_id'],source['id']}
    assert len(unit['evidence_trail'])==2


def test_tensions_visible_when_both_groups_have_same_sources(knowledge):
    from app.services.workshop import synthesis_context
    common={'source_ids':['source_a','source_b'],'asset_ids':['first','second']}
    left={**common,'canonical_asset':knowledge}
    right={**common,'canonical_asset':{**knowledge,'id':'opposing','what_it_says':'Do not use practice breaks to help participants retain learning.'}}
    assert synthesis_context([left,right])['tensions']

"""Probe existing pure parser functions on synthetic documents; no app startup."""
import ast
import base64
import hashlib
import json
from pathlib import Path
import socket
import tempfile
from io import BytesIO

from docx import Document
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, NumberObject, DecodedStreamObject

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]


def pure_functions(rel):
    path=ROOT/rel
    tree=ast.parse(path.read_text())
    # Execute exact existing definitions/imports only; exclude application state imports/initializers.
    nodes=[]
    for n in tree.body:
        if isinstance(n,(ast.FunctionDef,ast.ClassDef)):
            nodes.append(n)
        elif isinstance(n,ast.Import):
            nodes.append(n)
        elif isinstance(n,ast.ImportFrom) and not n.module.startswith('app.'):
            nodes.append(n)
    ns={}
    exec(compile(ast.Module(body=nodes,type_ignores=[]),str(path),'exec'),ns)
    return ns,hashlib.sha256(path.read_bytes()).hexdigest()


def run():
    parser,ph=pure_functions('apps/api/app/services/text_extraction.py')
    chunker,ch=pure_functions('apps/api/app/services/text_chunking.py')
    structure,sh=pure_functions('apps/api/app/services/source_structure.py')
    png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1sAAAAASUVORK5CYII=')
    results={}
    with tempfile.TemporaryDirectory() as td:
        d=Path(td)
        doc=Document();doc.add_heading('Synthetic station report',0)
        doc.add_paragraph('Before table: counts are not safety certifications.')
        t=doc.add_table(rows=3,cols=2)
        for row,values in zip(t.rows,[['Category','Count (items)'],['Inspected','12'],['Referred','3']]):
            for c,v in zip(row.cells,values):c.text=v
        doc.add_paragraph('Footnote: referrals are not completed repairs.')
        doc.add_picture(BytesIO(png));doc.add_paragraph('Figure 1. Synthetic colored indicator; colors are not described here.')
        doc.add_page_break();doc.add_paragraph('After page break: owner retains decision authority.')
        p=d/'synthetic.docx';doc.save(p);before=hashlib.sha256(p.read_bytes()).hexdigest()
        text=parser['extract_text_from_file'](str(p),'.docx')
        import zipfile
        with zipfile.ZipFile(p) as z:images=[n for n in z.namelist() if n.startswith('word/media/')]
        results['docx']={'original_unchanged':before==hashlib.sha256(p.read_bytes()).hexdigest(),
                         'embedded_images_in_original':len(images),'extracted_text':text,
                         'table_values_headers_footnote_retained':all(s in text for s in ['Category | Count (items)','Inspected | 12','Referred | 3','referrals are not completed repairs']),
                         'table_grid_in_plain_text':False,'image_payload_in_text':False,
                         'caption_retained':'Figure 1.' in text,'page_locators_in_text_contract':False,
                         'image_to_caption_relationship_in_text_contract':False}
        html='<h1>Diagram review</h1><p>Before image.</p><figure><img src="data:image/png;base64,'+base64.b64encode(png).decode()+'" alt="ALT_ONLY_RED"><figcaption>Figure 2. Visible caption.</figcaption></figure><p>After image.</p>'
        p=d/'synthetic.html';p.write_text(html)
        text=parser['extract_text_from_file'](str(p),'.html')
        results['html']={'extracted_text':text,'caption_retained':'Figure 2.' in text,
                         'image_alt_retained':'ALT_ONLY_RED' in text,'image_data_in_original':'data:image/png' in p.read_text(),
                         'image_association_in_text_contract':False}
        # PDF uses explicit positioned text and a vector diagram plus raster image on an image-only page.
        writer=PdfWriter()
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        fontref=writer._add_object(font)
        def page(content):
            pg=writer.add_blank_page(width=400,height=500)
            pg[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):fontref})})
            stream=DecodedStreamObject();stream.set_data(content.encode());pg[NameObject('/Contents')]=writer._add_object(stream)
            return pg
        page('BT /F1 12 Tf 20 460 Td (Category) Tj 180 0 Td (Count items) Tj -180 -30 Td (Inspected) Tj 180 0 Td (12) Tj -180 -30 Td (Referred) Tj 180 0 Td (3) Tj -180 -30 Td (Footnote: referrals are not repairs.) Tj ET')
        pg=page('1 0 0 rg 20 200 60 60 re f 0 0 1 rg 170 200 60 60 re f 0 0 0 RG 80 230 m 170 230 l S q 10 0 0 10 20 20 cm /Im0 Do Q')
        image=DecodedStreamObject();image.set_data(bytes([255,0,0]));image.update({NameObject('/Type'):NameObject('/XObject'),NameObject('/Subtype'):NameObject('/Image'),NameObject('/Width'):NumberObject(1),NameObject('/Height'):NumberObject(1),NameObject('/ColorSpace'):NameObject('/DeviceRGB'),NameObject('/BitsPerComponent'):NumberObject(8)})
        pg['/Resources'][NameObject('/XObject')]=DictionaryObject({NameObject('/Im0'):writer._add_object(image)})
        page('BT /F1 12 Tf 20 460 Td (Figure 3. Diagram appears on prior page.) Tj 0 -30 Td (Owner retains decision authority.) Tj ET')
        p=d/'synthetic.pdf'
        with p.open('wb') as f:writer.write(f)
        before=hashlib.sha256(p.read_bytes()).hexdigest();text=parser['extract_text_from_file'](str(p),'.pdf')
        chunks=chunker['chunk_text']('synthetic',text)
        chunks,annotation=structure['annotate_pdf_chunks_with_outline'](p,chunks)
        pages=[pg.extract_text() or '' for pg in PdfReader(p).pages]
        results['pdf']={'original_unchanged':before==hashlib.sha256(p.read_bytes()).hexdigest(),
                        'original_pages':len(pages),'image_only_page_retained_in_original':len(PdfReader(p).pages[1].images)>0,
                        'page_text_lengths':[len(x) for x in pages],'extracted_text':text,
                        'image_only_page_silently_absent_from_text':not pages[1].strip() and bool(text.strip()),
                        'table_structure_returned':False,'caption_retained':'Figure 3.' in text,
                        'chunk_page_starts':[b.get('source_page_start') for b in chunks],
                        'caption_diagram_link_returned':False,
                        'original_unchanged_after_annotation':before==hashlib.sha256(p.read_bytes()).hexdigest(),
                        'note':'Existing chunk metadata supplies a coarse starting page; individual extracted passages have no page map, grid or image relationship.'}
    return {'parser_hash':ph,'chunker_hash':ch,'structure_hash':sh,'fixtures':results,
            'method':'Exact pure function definitions loaded via AST to avoid app storage initialization. Tests parser/annotation behavior, not end-to-end upload persistence. Synthetic files generated in a disposable directory; no new dependencies.',
            'limitations':'One simple synthetic document per format; no OCR, diagram meaning or complex PDF parsing success is inferred.'}


if __name__=='__main__':
    def deny(*a,**k):raise AssertionError('No network permitted')
    socket.socket=deny;socket.create_connection=deny
    result=run();(HERE/'multimodal-results.json').write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result,indent=2))

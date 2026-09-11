import io
import json
import subprocess
import sys
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from docx import Document
from app.main import app
from app.services.text_chunking import chunk_text


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('KNOWLEDGE_ENGINE_STORAGE', str(tmp_path))
    with TestClient(app) as client:
        yield client


def source(client, content=b'Use small experiments to test uncertain assumptions before committing resources.', filename='notes.txt'):
    library = client.post('/libraries', json={'name': 'Decision intelligence'}).json()
    source = client.post('/sources', json={'library_id': library['id'], 'title': 'Experiment notes'}).json()
    assert client.post(f"/sources/{source['id']}/upload", files={'file': (filename, content)}).status_code == 200
    return source


def unit(client, sid):
    assert client.post(f'/sources/{sid}/process').status_code == 200
    chunk = client.get(f'/sources/{sid}/chunks').json()['items'][0]
    return {'chunk_id': chunk['id'], 'knowledge_type': 'principle', 'title': 'Test assumptions', 'insight': 'Use small experiments before committing resources.', 'evidence': 'Use small experiments to test uncertain assumptions'}


def test_complete_workflow_and_restart(client, tmp_path):
    s = source(client); sid=s['id']; u=unit(client,sid)
    assert u['chunk_id'] in client.get(f"/sources/{sid}/prompts/{u['chunk_id']}").text
    response=client.post(f'/sources/{sid}/knowledge', json={'units':[u]})
    assert response.status_code==201, response.text
    uid=response.json()['items'][0]['id']
    payload={'title':'Experiment plan','goal':'Validate the riskiest assumption','knowledge_ids':[uid]}
    assert client.post('/workshops',json=payload).status_code==422
    assert client.patch(f'/knowledge/{uid}/review',json={'status':'approved'}).status_code==200
    assert client.post(f'/sources/{sid}/knowledge',json={'units':[u]}).json()['items'][0]['status']=='approved'
    assert client.get('/knowledge').json()['count']==1
    w=client.post('/workshops',json=payload).json()
    assert 'Experiment notes' in w['content'] and u['evidence'] in w['content']
    assert client.patch(f"/workshops/{w['id']}",json={'content':'# My plan\nRun one experiment. [1]'}).status_code==200
    assert 'Run one experiment' in client.get(f"/workshops/{w['id']}/export").text
    # New process proves data is not only surviving inside module memory.
    code="from app.db import store; print(store.all_records('workshop')[0]['content'])"
    result=subprocess.run([sys.executable,'-c',code],capture_output=True,text=True,check=True)
    assert 'Run one experiment' in result.stdout
    assert client.delete(f'/sources/{sid}').status_code==409
    assert client.delete(f"/workshops/{w['id']}").status_code==204
    assert client.delete(f'/sources/{sid}').status_code==204
    assert client.get('/knowledge').json()['count']==0
    assert not list((tmp_path/'uploads').glob('*'))


def test_invalid_batch_is_atomic_and_traceable(client):
    sid=source(client)['id'];u=unit(client,sid)
    bad={**u,'title':'Invented claim','evidence':'This fabricated quotation is not in the source.'}
    assert client.post(f'/sources/{sid}/knowledge',json={'units':[u,bad]}).status_code==422
    assert client.get('/knowledge').json()['count']==0
    bad={**u,'chunk_id':'some-other-source'}
    assert client.post(f'/sources/{sid}/knowledge',json={'units':[bad]}).status_code==422


def test_input_validation_and_processing_errors(client):
    assert client.post('/libraries',json={'name':'   '}).status_code==422
    assert client.post('/sources',json={'title':'x','library_id':'missing'}).status_code==404
    sid=source(client, b' ', 'empty.txt')['id']
    assert client.post(f'/sources/{sid}/process').status_code==422
    assert client.get(f'/sources/{sid}/status').json()['processing_status']=='extraction_failed'
    sid=source(client, b'not-json', 'broken.json')['id']
    assert client.post(f'/sources/{sid}/process').status_code==400
    assert client.post(f'/sources/{sid}/upload',files={'file':('different.txt',b'replacement')}).status_code==409


def test_docx_extraction_includes_tables(client):
    doc=Document();doc.add_paragraph('An experiment makes an assumption testable.')
    table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Measure';table.cell(0,1).text='Demand'
    data=io.BytesIO();doc.save(data)
    sid=source(client,data.getvalue(),'source.docx')['id']
    assert client.post(f'/sources/{sid}/process').status_code==200
    text=client.get(f'/sources/{sid}/extracted-text').text
    assert 'assumption testable' in text and 'Measure | Demand' in text


def test_text_formats(client):
    for filename,content,expected in [('x.html',b'<p>Useful idea</p><script>ignore me</script>','Useful idea'),('x.csv',b'Rule,Reason\nTest,Uncertainty','Test | Uncertainty'),('x.json',b'{"principle":"Test first"}','Test first')]:
        sid=source(client,content,filename)['id']
        assert client.post(f'/sources/{sid}/process').status_code==200
        assert expected in client.get(f'/sources/{sid}/extracted-text').text


def test_chunk_coverage_without_redundant_tail():
    text=' '.join(str(i) for i in range(1100))
    chunks=chunk_text('x',text)
    assert len(chunks)==1
    assert chunks[0]['text']==text
    chunks=chunk_text('x',' '.join(str(i) for i in range(2300)))
    assert chunks[-1]['end_word_index']==2300
    for left,right in zip(chunks,chunks[1:]):
        assert left['end_word_index']-right['start_word_index']==200
    with pytest.raises(ValueError):chunk_text('x','text',10,-1)


def test_ui_assets_and_cross_origin_writes(client):
    assert client.get('/').status_code==200
    assert 'Knowledge Engine' in client.get('/').text
    for file in ['app.js','style.css','icon.svg']:
        assert client.get('/workspace/'+file).status_code==200
    assert client.post('/libraries',json={'name':'x'},headers={'Origin':'https://foreign.example'}).status_code==403
